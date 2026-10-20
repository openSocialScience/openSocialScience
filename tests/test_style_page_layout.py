"""风格包的页面级版式：页边距、页眉页脚、页码、单双面、分栏、行号、脚注编号。

**这一层原先整个不存在。**段落级档位有 `_ZH_LAYOUT` 那张表，页面级的东西全在
docx 的 `sectPr` 与 `settings.xml` 里，风格包只能手工塞一个 `template.docx`——
于是「这本刊要奇偶页不同页眉」既没法声明也没法校验。

建三个中文刊的包时顺带逼出五处缺陷，一并在这里钉住：
中文数量单位（「2万字」）、中文引用说法的落回、`csl_for` 把中文夹注判成 Chicago、
摘要字数两个口径、以及 pandoc 丢掉 `footnotePr`。
"""

from __future__ import annotations

import shutil
import zipfile

import pytest

from oss.core.contracts import PageLayout, StyleRules
from oss.core.style import _numbers_in, _style_word_in, check_rules_grounded, csl_for

pytest.importorskip("docx")
PANDOC = shutil.which("pandoc")
pytestmark = pytest.mark.skipif(not PANDOC, reason="没装 pandoc")


def _built(tmp_path, page: PageLayout):
    """按给定版式建一份模板，返回打开后的 Document 与它的 section。"""
    import docx

    from oss.api.paper import _ZH_LAYOUT, _build_reference_doc

    out = _build_reference_doc(PANDOC, tmp_path / "t.docx", _ZH_LAYOUT, "Times New Roman",
                               page=page, fields={"journal": "《社会学研究》", "title": "某题",
                                                  "year": "2026"},
                               stamp=str(abs(hash(page.model_dump_json())))[:8])
    assert out is not None
    d = docx.Document(str(out))
    return d, d.sections[0]


def test_页边距落到毫米级(tmp_path) -> None:
    _, sec = _built(tmp_path, PageLayout(margin_top_cm=2.5, margin_bottom_cm=2.5,
                                         margin_left_cm=3.0, margin_right_cm=2.0))
    got = [round(getattr(sec, a).cm, 2) for a in
           ("top_margin", "bottom_margin", "left_margin", "right_margin")]
    assert got == [2.5, 2.5, 3.0, 2.0]


def test_单双面要同时开镜像页边距与奇偶页页眉(tmp_path) -> None:
    """**只设 even_page_header 不开 evenAndOddHeaders，偶数页页眉根本不显示。**
    两个开关都是文档级的，不在 sectPr 里。
    """
    from docx.oxml.ns import qn

    d, sec = _built(tmp_path, PageLayout(two_sided=True, header_odd="{journal}",
                                         header_even="{title}"))
    st = d.settings.element
    assert st.find(qn("w:mirrorMargins")) is not None
    assert st.find(qn("w:evenAndOddHeaders")) is not None
    assert sec.header.paragraphs[0].text == "《社会学研究》"
    assert sec.even_page_header.paragraphs[0].text == "某题"


def test_单面时不开镜像页边距(tmp_path) -> None:
    from docx.oxml.ns import qn

    d, _ = _built(tmp_path, PageLayout(margin_left_cm=3.0))
    assert d.settings.element.find(qn("w:mirrorMargins")) is None


def test_页码是真的域不是死数字(tmp_path) -> None:
    """写「1」出来就是每页都印 1。页码必须是 PAGE 域，Word 打开时才按页重算。"""
    d, sec = _built(tmp_path, PageLayout(page_number="footer-center", page_number_start=1))
    xml = sec.footer.paragraphs[0]._p.xml
    assert "PAGE" in xml and "fldChar" in xml


def test_页眉已有文字时页码另起一段_不冲掉文字(tmp_path) -> None:
    d, sec = _built(tmp_path, PageLayout(header_odd="{journal}", page_number="header-right"))
    texts = [p.text for p in sec.header.paragraphs]
    assert "《社会学研究》" in texts, texts


@pytest.mark.parametrize("fmt,want", [("arabic", "decimal"), ("roman-lower", "lowerRoman"),
                                      ("roman-upper", "upperRoman")])
def test_页码格式(tmp_path, fmt: str, want: str) -> None:
    from docx.oxml.ns import qn

    _, sec = _built(tmp_path, PageLayout(page_number="footer-center", page_number_format=fmt))
    assert sec._sectPr.find(qn("w:pgNumType")).get(qn("w:fmt")) == want


def test_分栏与行号(tmp_path) -> None:
    from docx.oxml.ns import qn

    _, sec = _built(tmp_path, PageLayout(columns=2, column_gap_cm=0.8, line_numbers=True))
    cols = sec._sectPr.find(qn("w:cols"))
    assert cols.get(qn("w:num")) == "2" and cols.get(qn("w:space")) == "453"
    assert sec._sectPr.find(qn("w:lnNumType")).get(qn("w:countBy")) == "1"


def test_一栏时不写cols(tmp_path) -> None:
    """缺省一栏。写 `w:cols num=1` 不算错，但**不动它**才能让「须知没说就不设」名副其实。"""
    from docx.oxml.ns import qn

    _, sec = _built(tmp_path, PageLayout(columns=1))
    cols = sec._sectPr.find(qn("w:cols"))
    assert cols is None or cols.get(qn("w:num")) in (None, "1")


def test_pandoc丢掉的脚注编号要在成稿里补回(tmp_path) -> None:
    """**实测：pandoc 保留页边距、镜像页边距、奇偶页页眉、分栏、页码类型、行号、
    页眉页脚部件，唯独 settings.xml 里的 `w:footnotePr` 是它自己重写的。**
    于是模板里声明好的「每页单独编号」在成稿里变回 Word 缺省的连续编号——
    声明了却不生效，而且要翻到第二页去数脚注号才发现。
    """
    from oss.api.paper import _patch_footnote_restart

    src = tmp_path / "out.docx"
    with zipfile.ZipFile(src, "w") as z:
        z.writestr("word/settings.xml",
                   '<?xml version="1.0"?><w:settings xmlns:w="http://schemas.openxmlformats.org'
                   '/wordprocessingml/2006/main"><w:zoom/></w:settings>')
        z.writestr("word/document.xml", "<x/>")
    _patch_footnote_restart(src)
    with zipfile.ZipFile(src) as z:
        s = z.read("word/settings.xml").decode()
        assert "eachPage" in s and "footnotePr" in s
        assert "word/document.xml" in z.namelist()   # 重打包不许丢部件
    _patch_footnote_restart(src)                     # 幂等：跑两遍不该叠出两个
    with zipfile.ZipFile(src) as z:
        assert z.read("word/settings.xml").decode().count("footnotePr") == 2   # 开合各一


# ── 建中文风格包时逼出来的五处 ────────────────────────────────────
def test_中文的数量单位要展开() -> None:
    """**中文期刊的须知几乎不写阿拉伯数字的整数**——写的是「不超 2 万字」。
    只抓阿拉伯数字的话，`word_limit: 20000` 这个**填对了**的值会被判成
    「须知里没有这个数」，于是每一个正确的中文包都过不了那道门。
    """
    assert "20000" in _numbers_in("单篇稿件字数不超2万字。")
    assert "12000" in _numbers_in("以1.2万—2万字为宜")
    assert "20000" in _numbers_in("以1.2万—2万字为宜")
    assert "2" in _numbers_in("不超2万字")          # 原样的数仍然认


def test_中文的引用说法也要落回须知() -> None:
    """原先取 `split()[0]`——英文假设。中文串没有空格，整串被当成一个词去找。"""
    assert _style_word_in("页下注（每页单独编号）", "注释请用页下注（每页单独编号）")
    assert _style_word_in("Chicago Manual of Style, 17th edition", "... Chicago Manual of Style ...")
    assert not _style_word_in("温哥华制", "注释请用页下注")


@pytest.mark.parametrize("said,want", [
    # **中文刊拿到英文体例**是原先的行为：这句白话不含「著者」「顺序编码」等术语，
    # 于是落到最后的 Chicago 兜底上。更坏的是不带风格包时反而给的是国标。
    ("在引文后加括号注明作者、出版年份及页码", "china-national-standard-gb-t-7714-2015-author-date"),
    ("注释请用页下注（每页单独编号）", "china-national-standard-gb-t-7714-2015-note"),
    ("参考文献采用顺序编码制，注释的序号为连续编号", "china-national-standard-gb-t-7714-2015-numeric"),
    ("Chicago Manual of Style, 17th edition", "chicago-author-date"),
    ("APA 7th", "apa"),
])
def test_引用体例按须知的说法判(said: str, want: str) -> None:
    assert csl_for(said) == want


def test_版式里的数字也要落回须知原文() -> None:
    """**版式也是须知的一部分。**按印象填一个页边距，跟编一个字数上限是同一件事，
    而且更隐蔽：读者打开 Word 看到的就是这个版式，会以为它来自该刊的要求。
    """
    guide = "本刊要求页边距上下各 2.5 厘米，正文一栏。"
    assert check_rules_grounded(StyleRules(journal="X", layout=PageLayout(margin_top_cm=2.5)), guide) == []
    bad = check_rules_grounded(StyleRules(journal="X", layout=PageLayout(margin_left_cm=3.7)), guide)
    assert bad and "layout.margin_left_cm" in bad[0]
    assert check_rules_grounded(StyleRules(journal="X"), guide) == []   # 没声明版式就不判


def test_摘要字数两处口径要一致() -> None:
    """一处量 298 字、一处量 334 字，都叫「摘要字数」——差的是**取法**不是数法：
    按节切会把「## 摘要」标题与「**关键词**：……」那一行算进去。
    """
    from oss.core.manuscript import _abstract_body
    from oss.core.style import check_paper, count_words

    paper = ("# 题目\n\n## 摘要\n\n" + "信" * 210 + "。\n\n"
             "**关键词**：甲；乙；丙\n\n## 一、引言\n\n正文。\n")
    assert count_words(_abstract_body(paper).strip()) == 210   # 210 个「信」，句号与关键词行不算
    problems, _ = check_paper(paper, StyleRules(journal="X", abstract_max_words=220), None, {})
    assert not any("摘要" in p for p in problems), problems


def test_没有作者的记录印占位_并在导出前就警告() -> None:
    """**「Anonymous」不是占位，那是一个像样的作者名。**真跑时成稿正文印的是
    「(Anonymous; Delhey 等, 2011)」，读者与作者都会以为那是一篇佚名著作。
    题名那一侧的约定是 `[UNVERIFIED]`，作者这一侧照同一形状。
    """
    from oss.core.contracts import LiteratureCorpus, PaperRecord
    from oss.core.manuscript import NO_AUTHOR, _no_author_warnings, _split_name

    assert _split_name("") == {"family": NO_AUTHOR}
    assert NO_AUTHOR.startswith("[") and NO_AUTHOR.endswith("]")   # 方括号是关键：一眼看得出是占位
    corpus = LiteratureCorpus(records=[
        PaperRecord(paper_id="pdf_x", title="[UNVERIFIED] 某文", authors=[]),
        PaperRecord(paper_id="pdf_y", title="有作者的", authors=["张三"]),
    ])
    w = _no_author_warnings("正文 {{cite:pdf_x}} 与 {{cite:pdf_y}}。", corpus)
    assert len(w) == 1 and "pdf_x" in str(w[0]) and "pdf_y" not in str(w[0])
    assert _no_author_warnings("只引有作者的 {{cite:pdf_y}}。", corpus) == []


# ── 注文的字体字号：须知写了，原先没人读 ──────────────────────────
@pytest.mark.parametrize("said,want_pt", [
    ("五号", 10.5), ("5号", 10.5), ("小五", 9), ("10.5pt", 10.5), ("12", 12.0),
    ("三号", 16), ("说不清的字号", None), ("", None),
])
def test_中文字号名换算(said: str, want_pt: float | None) -> None:
    """须知写的是字号名（「注文排5号仿宋体」），**照抄那个名字**比自己折成磅安全：
    五号＝10.5 磅是国标里的固定对应，工具查表换算；折算这一步一旦写错，
    成品上看不出来，而读者会以为那就是该刊的要求。
    """
    from oss.api.paper import _font_pt

    assert _font_pt(said) == want_pt


def test_注文字体字号真的落到样式上(tmp_path) -> None:
    """**这条要求原先只躺在 `other_requirements` 里**——那是给写作的人看的散文，
    没有任何代码读它。于是《中国社会科学》须知白纸黑字写着「注文排5号仿宋体」，
    导出来的却是缺省的小五（9 磅），差一档，而成品上看不出来。
    """
    from docx.shared import Pt

    from oss.core.contracts import PageLayout

    d, _ = _built(tmp_path, PageLayout(footnote_font="仿宋", footnote_size="5号"))
    st = next(x for x in d.styles if x.name == "Footnote Text")
    assert st.font.size == Pt(10.5)
    from docx.oxml.ns import qn

    assert st.element.get_or_add_rPr().get_or_add_rFonts().get(qn("w:eastAsia")) == "仿宋"


def test_注文的字体字号也要落回须知原文() -> None:
    """会真的改成品排版的每一格都要有出处。按印象填一个「小五仿宋」，
    读者打开 Word 看到的就是它，会以为那是该刊的要求。"""
    from oss.core.contracts import PageLayout, StyleRules

    guide = "注释请用页下注（每页单独编号），注文排5号仿宋体。"
    ok = StyleRules(journal="X", layout=PageLayout(footnote_font="仿宋", footnote_size="5号"))
    assert check_rules_grounded(ok, guide) == []
    bad = StyleRules(journal="X", layout=PageLayout(footnote_size="小五"))
    assert check_rules_grounded(bad, guide)


def test_纸张与正文字体字号也落到样式上(tmp_path) -> None:
    """《中国社会科学》须知：「稿件请用A4纸格式……研究论文类稿件宋体小4号」。
    跟注文那一处同一个毛病——原先只抄在 `other_requirements` 里，没有代码读它。
    pandoc 模板缺省是 Letter 纸，中文期刊一律 A4。
    """
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    from oss.core.contracts import PageLayout

    d, sec = _built(tmp_path, PageLayout(paper_size="A4", body_font="宋体", body_size="小4号"))
    assert abs(sec.page_width - Mm(210)) < Mm(1) and abs(sec.page_height - Mm(297)) < Mm(1)
    st = next(x for x in d.styles if x.name == "Normal")
    assert st.font.size == Pt(12)
    assert st.element.get_or_add_rPr().get_or_add_rFonts().get(qn("w:eastAsia")) == "宋体"
    # 首行缩进两字符要跟着字号走：写死的磅数换了字号就不是两个字了
    assert st.paragraph_format.first_line_indent == Pt(24)


def test_纸张没说就不动(tmp_path) -> None:
    """须知没规定纸张时不该擅自改——「留空就不动」是这一整块的规矩。"""
    from oss.core.contracts import PageLayout

    _, a = _built(tmp_path, PageLayout(margin_top_cm=2.5))
    _, b = _built(tmp_path, PageLayout(margin_top_cm=2.5, paper_size="A5"))
    assert a.page_width != b.page_width


@pytest.mark.parametrize("field,val", [
    ("paper_size", "A5"), ("body_font", "楷体"), ("body_size", "三号"),
])
def test_版式里会改排版的每一格都要落回须知(field: str, val: str) -> None:
    from oss.core.contracts import PageLayout, StyleRules

    guide = "稿件请用A4纸格式，正文简体横排，研究论文类稿件宋体小4号。"
    bad = check_rules_grounded(StyleRules(journal="X", layout=PageLayout(**{field: val})), guide)
    assert bad and field in bad[0], f"{field}={val} 该被拦下"


# ── 字数：中文期刊的要求常常是区间，只管上限等于只管了一半 ──────────
def _paper(n_body: int, n_refs: int = 0) -> str:
    refs = ("\n\n## 参考文献\n\n" + "参" * n_refs) if n_refs else ""
    return "# 题目\n\n## 摘要\n\n摘要。\n\n## 一、引言\n\n" + "字" * n_body + refs


def test_低于下限要出警告_但不拦导出() -> None:
    """**须知那句是「以 X—Y 字为宜」，是建议不是硬杠**，所以只警告；
    但漏掉它等于只管了一半——写得太短一样退稿。
    真跑时描述那篇 6852 字，而《社会学研究》要一万到两万。
    """
    from oss.core.style import check_paper

    r = StyleRules(journal="社会学研究", word_min=10000, word_limit=20000)
    probs, warns = check_paper(_paper(6000), r, None, {})
    assert not any("低于" in p for p in probs), "下限不该拦导出"
    assert any("低于" in w and "10000" in w for w in warns), warns
    # 在区间里就什么都不说
    _, w2 = check_paper(_paper(12000), r, None, {})
    assert not any("低于" in x or "超过" in x for x in w2)


def test_超过上限仍然拦() -> None:
    from oss.core.style import check_paper

    r = StyleRules(journal="X", word_limit=1000)
    probs, _ = check_paper(_paper(2000), r, None, {})
    assert any("超过" in p for p in probs)


def test_含不含参考文献要真的算进去_不只写在提示里() -> None:
    """`word_limit_includes_references` 原先**只用在提示文字里，没用在计算上**——
    不管须知怎么说，算的都是全文，于是「上限不含参考文献」的刊被判得过严。
    """
    from oss.core.style import check_paper

    text = _paper(900, n_refs=400)          # 正文 900，参考文献 400
    inc = StyleRules(journal="X", word_limit=1000, word_limit_includes_references=True)
    exc = StyleRules(journal="X", word_limit=1000, word_limit_includes_references=False)
    assert any("超过" in p for p in check_paper(text, inc, None, {})[0]), "含参考文献该超"
    assert not any("超过" in p for p in check_paper(text, exc, None, {})[0]), "不含参考文献不该超"


def test_下限也要落回须知原文() -> None:
    guide = "学术论文以10000-20000字为宜。"
    ok = StyleRules(journal="X", word_min=10000, word_limit=20000)
    assert check_rules_grounded(ok, guide) == []
    bad = StyleRules(journal="X", word_min=8000)
    assert check_rules_grounded(bad, guide)
