"""中文 Word 导出（第 4 步）：宋体模板、GB/T 7714 参考文献。

- 缺省样式跟稿子语言走：英文 Chicago 作者-年份，中文 GB/T 7714 **著者-出版年**（夹注）
- 中文模板**按档分开**：正文、摘要、表格内文、题注、表注、参考文献各一档，不再一刀切
- 表是三线表：顶线、栏目线、底线，没有竖线
- 国标 CSL 真的能被 citeproc 用起来：正文出编号 / 夹注，书目按国标排
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest

from oss.api.paper import default_csl_name, default_reference_doc_zh
from oss.api.paper import csl_path

pandoc = shutil.which("pandoc")


def test_缺省参考文献样式跟稿子语言走() -> None:
    """中文缺省是**夹注**不是顺序编码——社会学期刊那一路（2026-09-16 用户定的「甲」）。"""
    assert default_csl_name("## 摘要\n本文考察教育扩张。") == "china-national-standard-gb-t-7714-2015-author-date"
    assert default_csl_name("## Abstract\nWe study expansion.") == "chicago-author-date"


def test_两个国标CSL在包里() -> None:
    for name in ("china-national-standard-gb-t-7714-2015-numeric",
                 "china-national-standard-gb-t-7714-2015-author-date"):
        p = csl_path(name)
        assert p is not None and p.exists(), name
        assert "GB/T 7714" in p.read_text(encoding="utf-8")[:2000]


def _zh_styles(tmp_path: Path, monkeypatch):
    """建一份中文模板，返回 {样式名: 样式}。缓存写进临时目录的 styles/_default/。"""
    pytest.importorskip("docx")
    monkeypatch.chdir(tmp_path)
    doc = default_reference_doc_zh(pandoc)
    assert doc is not None and doc.exists()
    import docx

    return doc, {s.name: s for s in docx.Document(str(doc)).styles}


@pytest.mark.skipif(pandoc is None, reason="没装 pandoc")
def test_中文模板_宋体黑体_行距_首行缩进(tmp_path: Path, monkeypatch) -> None:
    doc, _ = _zh_styles(tmp_path, monkeypatch)
    styles = zipfile.ZipFile(doc).read("word/styles.xml").decode("utf-8")
    assert 'w:eastAsia="宋体"' in styles
    assert 'w:eastAsia="黑体"' in styles, "标题用黑体"
    assert "Times New Roman" in styles, "拉丁字符保留西文字体"
    assert 'w:firstLine="480"' in styles, "正文首行缩进两字符（24pt = 480 twips）"
    assert default_reference_doc_zh(pandoc) == doc, "第二次用缓存"


@pytest.mark.skipif(pandoc is None, reason="没装 pandoc")
def test_各档分开_摘要表格题注参考文献不跟正文一个样(tmp_path: Path, monkeypatch) -> None:
    """用户在成稿上抓到的：「摘要关键词正文图下注解都用的一个格式」。

    根因是一个 `_BODY_LIKE` 集合把十来个样式刷成同一套字号行距。这里逐档回验。
    """
    from docx.shared import Pt, RGBColor

    _, st = _zh_styles(tmp_path, monkeypatch)

    body = st["Normal"]
    assert body.font.size == Pt(12) and body.paragraph_format.line_spacing == 1.5
    assert body.paragraph_format.first_line_indent == Pt(24)

    cell = st["Compact"]                       # pandoc 给每个表格单元格挂这一档
    assert cell.font.size == Pt(10.5), "表格内文要比正文小"
    assert cell.paragraph_format.line_spacing == 1.0, "表格内文不跟正文一样 1.5 倍行距"
    assert cell.paragraph_format.first_line_indent == Pt(0), "表格内文不能继承正文的首行缩进"

    bib = st["Bibliography"]
    assert bib.paragraph_format.first_line_indent == Pt(-21), "参考文献是悬挂缩进，不是首行缩进"
    assert bib.paragraph_format.left_indent == Pt(21)

    for name in ("Caption", "Image Caption", "Table Caption"):
        assert st[name].font.italic is False, f"{name}：汉字没有真正的斜体"
        assert st[name].font.size == Pt(10.5), name

    assert st["Abstract"].font.size == Pt(10.5) and st["Abstract"].paragraph_format.right_indent == Pt(24)
    assert st["Keywords"] is not None, "pandoc 的 custom-style 按样式名挂靠，模板里得有这一档"
    assert st["Table Note"].font.size == Pt(9), "表注是全篇最小的一档"

    assert st["Heading 2"].font.color.rgb == RGBColor.from_string("000000"), \
        "各级标题不能是 Word 主题的强调蓝（成稿里肉眼可见）"


@pytest.mark.skipif(pandoc is None, reason="没装 pandoc")
def test_表是三线表_有顶线底线没有竖线(tmp_path: Path, monkeypatch) -> None:
    doc, _ = _zh_styles(tmp_path, monkeypatch)
    styles = zipfile.ZipFile(doc).read("word/styles.xml").decode("utf-8")
    table = re.search(r'<w:style [^>]*w:styleId="Table"[^>]*>.*?</w:style>', styles, re.S).group(0)
    borders = re.search(r"<w:tblBorders>.*?</w:tblBorders>", table, re.S)
    assert borders is not None, "原样式只有一条栏目线，顶线底线都没有"
    assert "<w:top " in borders.group(0) and "<w:bottom " in borders.group(0), "顶线与底线"
    assert "<w:left " not in borders.group(0) and "<w:right " not in borders.group(0), "三线表没有竖线"
    assert 'w:type="firstRow"' in table, "栏目线挂在表头行下面"


@pytest.mark.skipif(pandoc is None, reason="没装 pandoc")
def test_三线表不用后处理_转出来的文档就带着(tmp_path: Path, monkeypatch) -> None:
    """pandoc 写出来的表只挂 `<w:tblStyle w:val="Table"/>`、不自带边框，
    所以改样式表就生效。这一条一旦不成立，三线表就得靠后处理去补。"""
    doc, _ = _zh_styles(tmp_path, monkeypatch)
    (tmp_path / "t.md").write_text("| 变量 | 估计值 |\n|---|---:|\n| 阶级 | -0.140 |\n", encoding="utf-8")
    r = subprocess.run([pandoc, "t.md", "-f", "markdown", "--reference-doc", str(Path(doc).resolve()),
                        "-o", "t.docx"], cwd=tmp_path, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    with zipfile.ZipFile(tmp_path / "t.docx") as z:
        body = z.read("word/document.xml").decode("utf-8")
        styles = z.read("word/styles.xml").decode("utf-8")
    assert "<w:tblBorders>" not in body, "文档自带边框的话，样式表就压不住它"
    table = re.search(r'<w:style [^>]*w:styleId="Table"[^>]*>.*?</w:style>', styles, re.S).group(0)
    assert "<w:tblBorders>" in table, "三线表跟着模板进了转出来的文档"


@pytest.mark.skipif(pandoc is None, reason="没装 pandoc")
def test_档位改版后旧缓存会重建(tmp_path: Path, monkeypatch) -> None:
    """不认版本的话，改完排版的人本地那份旧缓存会一直用下去，怎么改都看不出变化。"""
    import docx

    from oss.api import paper as paper_api

    doc, _ = _zh_styles(tmp_path, monkeypatch)
    assert docx.Document(str(doc)).core_properties.comments == paper_api._LAYOUT_VERSION
    monkeypatch.setattr(paper_api, "_LAYOUT_VERSION", "999")
    default_reference_doc_zh(pandoc)
    assert docx.Document(str(doc)).core_properties.comments == "999", "版本对不上要重建"


@pytest.mark.skipif(pandoc is None, reason="没装 pandoc")
def test_国标顺序编码真的排出编号与书目(tmp_path: Path) -> None:
    (tmp_path / "refs.json").write_text(json.dumps([{
        "id": "zhang2020", "type": "article-journal",
        "title": "教育扩张与地位获得", "container-title": "某学刊",
        "author": [{"family": "张", "given": "三"}], "issued": {"date-parts": [[2020]]},
        "volume": "35", "issue": "2", "page": "1-20",
    }], ensure_ascii=False), encoding="utf-8")
    (tmp_path / "p.md").write_text(
        "已有研究考察了这一问题 [@zhang2020]。\n\n# 参考文献\n\n::: {#refs}\n:::\n", encoding="utf-8")
    csl = csl_path("china-national-standard-gb-t-7714-2015-numeric")
    r = subprocess.run([pandoc, "p.md", "-f", "markdown", "--citeproc",
                        "--csl", str(csl), "--bibliography", "refs.json", "-t", "plain"],
                       cwd=tmp_path, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    assert "[1]" in r.stdout, r.stdout
    assert "张三" in r.stdout and "教育扩张与地位获得" in r.stdout, r.stdout

    r2 = subprocess.run([pandoc, "p.md", "-f", "markdown", "--citeproc",
                         "--csl", str(csl), "--bibliography", "refs.json", "-o", "out.docx"],
                        cwd=tmp_path, capture_output=True, text=True)
    assert r2.returncode == 0 and (tmp_path / "out.docx").exists(), r2.stderr


# ══════════════════════════════ 编号体例下叙述式引用不许悬空


def test_判断哪些样式是编号体例() -> None:
    from oss.api.paper import _is_numeric_csl

    assert _is_numeric_csl("china-national-standard-gb-t-7714-2015-numeric")
    assert not _is_numeric_csl("china-national-standard-gb-t-7714-2015-author-date")
    assert not _is_numeric_csl("chicago-author-date")
    assert _is_numeric_csl("vancouver")
    assert not _is_numeric_csl("")


@pytest.mark.skipif(pandoc is None, reason="没装 pandoc")
def test_国标编号下叙述式引用排出作者加上标而不是裸编号(tmp_path: Path) -> None:
    """验收跑 r20260907_edufert_zh 的 Word 里 33 处引用有 25 处长这样：

        「……几乎成为常识。[1] 把它写成了一个明确的理论命题……」

    一个上标编号当句子主语，中文里不成句；GB/T 7714 的叙述式引用要求作者名在正文里。
    这里从渲染器的输出一路过 citeproc，钉住最终排出来的是「张三[1] 指出」。
    """
    (tmp_path / "refs.json").write_text(json.dumps([{
        "id": "ref_a", "type": "article-journal", "title": "教育扩张与地位获得",
        "container-title": "某学刊", "author": [{"family": "张", "given": "三"}],
        "issued": {"date-parts": [[2020]]}, "volume": "35", "page": "1-20",
    }], ensure_ascii=False), encoding="utf-8")
    # 渲染器在编号体例下产出的形状：作者名 + 抑制作者的引用
    (tmp_path / "p.md").write_text(
        "张三[-@ref_a] 把它写成了一个明确的理论命题。\n\n# 参考文献\n\n::: {#refs}\n:::\n",
        encoding="utf-8")
    csl = csl_path("china-national-standard-gb-t-7714-2015-numeric")
    r = subprocess.run([pandoc, "p.md", "-f", "markdown", "--citeproc",
                        "--csl", str(csl), "--bibliography", "refs.json", "-t", "plain"],
                       cwd=tmp_path, capture_output=True, text=True)
    assert r.returncode == 0, r.stderr
    body = r.stdout.split("参考文献")[0]
    # `-t plain` 把上标写成 `^(...)`；Word 里是真上标。要钉的是**编号紧跟在作者名后面**
    assert re.search(r"张三\s*\^?\(?\[1\]", body), r.stdout
    assert not body.strip().startswith("[1]"), "编号不许当句子主语"
